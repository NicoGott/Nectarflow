#!/usr/bin/env python3
#
# Compares a replica taxon table with the published one (the Dryad workbook of
# Lowe et al. 2022). Their sheets carry consensus names built from both markers,
# ours single-marker ones, so each of our names is translated through their own
# 'consensus_identifications' sheet before anything is compared.
# What the numbers mean: docs/replica.md#9-compare-with-the-publication-compare
#
import argparse, csv, re, statistics, collections


#
# Puts a sample code into one spelling: the workbook writes P_007 where the
# samplesheet writes P_7
def norm(code):
    return re.sub(r"_0*(\d+)$", r"_\1", str(code).strip())


#
# The genera a taxon label refers to
# 'Cirsium/Centaurea spp.' -> {'Cirsium','Centaurea'}
# 'Centaurea nigra'        -> {'Centaurea'}
def genera(label):
    label = label.replace(" spp.", "").strip()
    if "/" in label:
        return {g.strip() for g in label.split("/")}
    return {label.split()[0]}


#
# Tie-averaged ranks of a list, in its own order
def ranks(v):
    order = sorted(range(len(v)), key=lambda i: v[i])
    out, i = [0.0] * len(v), 0
    while i < len(order):                   # walk each run of equal values together
        j = i
        while j + 1 < len(order) and v[order[j + 1]] == v[order[i]]:
            j += 1
        for k in range(i, j + 1):           # and give the whole run their mean rank
            out[order[k]] = (i + j) / 2 + 1
        i = j + 1
    return out


#
# Spearman's rho, i.e. Pearson on the ranks. statistics.correlation can rank for us with
# method="ranked", but that is Python 3.12 and the pipeline's environment is on 3.11.
# It raises for fewer than two points or a constant side, where nan is the honest answer
def spearman(x, y):
    if len(x) < 3:
        return float("nan")
    try:
        return statistics.correlation(ranks(x), ranks(y))
    except statistics.StatisticsError:
        return float("nan")


#
# Reads both tables, translates our taxon names onto theirs, then writes a
# per-sample comparison and the name translation it used
def main():
    ap = argparse.ArgumentParser(description="Compares a replica taxon table with the published one",
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--replica", required=True)
    ap.add_argument("--published", required=True)
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--marker", default=None, help="marker name in the consensus sheet (rbcL / ITS2); default: from --sheet")
    ap.add_argument("--meta-cols", type=int, default=8)
    ap.add_argument("--taxonomy", default=None, help="reference taxonomy TSV with a lineage column, to climb to tribe/family")
    ap.add_argument("--rank", choices=("label", "family"), default="label",
                    help="compare the published labels, or roll both sides up to family first")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    import openpyxl
    marker = a.marker or a.sheet.split("_")[0]
    wb = openpyxl.load_workbook(a.published, read_only=True)

    # row 1 holds the taxon names, the samples start at row 2
    rows = list(wb[a.sheet].iter_rows(values_only=True))
    taxa_pub = [str(t).strip() for t in rows[1][a.meta_cols:]]
    pub = {norm(r[0]): {t: float(v or 0) for t, v in zip(taxa_pub, r[a.meta_cols:])} for r in rows[2:] if r[0]}

    # genus -> published consensus labels, from their mapping table and from the sheet's own labels
    g2c = collections.defaultdict(set)
    for taxon, mk, cons in list(wb["consensus_identifications"].iter_rows(values_only=True))[2:]:
        if taxon and cons and str(mk).strip() == marker:
            for g in genera(str(taxon)):
                g2c[g].add(str(cons).strip())
    for c in taxa_pub:
        for g in genera(c):
            g2c[g].add(c)

    rep_rows = list(csv.reader(open(a.replica), delimiter="\t"))
    samples = [norm(s) for s in rep_rows[0][1:]]
    pub_total = collections.Counter()
    for d in pub.values():
        pub_total.update(d)
    labels = set(taxa_pub)
    lineage = {}                                    # genus -> list of ranks above it, nearest first
    if a.taxonomy:
        for r in csv.DictReader(open(a.taxonomy), delimiter="\t"):
            if r.get("lineage"):
                lineage.setdefault(r["genus"], list(reversed(r["lineage"].split(";"))))

    #
    # Maps one of our single-marker names onto a published label. Three tries,
    # in order, each one looser than the last:
    # 1. their sheet already has that exact label
    # 2. exactly one of their labels shares a genus with ours; if several do,
    #    the one carrying the most reads across their data wins
    # 3. climb our genus's lineage, nearest rank first, and take the first rank
    #    that is itself a label of theirs (they report some taxa no lower than
    #    tribe or family)
    # If none of the three fires the name is kept unchanged and reported as
    # having no counterpart: its reads still count against their sample total,
    # so a failed translation costs similarity rather than hiding
    def translate(name):
        if name in labels:                                             # 1. identical label
            return name, "identical"
        hits = set().union(*[g2c.get(g, set()) for g in genera(name)]) & labels
        if len(hits) == 1:                                             # 2. one label shares the genus
            return hits.pop(), "genus"
        if hits:                                                       # several: the most common in their data
            return max(hits, key=lambda h: pub_total[h]), "genus (ambiguous, most common chosen)"
        for g in sorted(genera(name)):                                 # 3. climb to tribe / family / order
            for rank in lineage.get(g, []):
                if rank in labels:
                    return rank, "higher rank"
        return name, "no published counterpart"

    family = {}                                     # a family, or a name under one -> that family
    if a.rank == "family":
        for r in csv.DictReader(open(a.taxonomy), delimiter="\t"):
            ranks = (r.get("lineage") or "").split(";")
            fam = next((x for x in ranks if x.endswith("aceae")), None)
            if fam:
                family.setdefault(r["genus"], fam)
                for x in ranks[ranks.index(fam):]:  # only the family and what sits under it: an
                    family.setdefault(x, fam)       # order holds many, and belongs to none of them

    #
    # The family a label belongs to, or the label itself when --rank is not
    # family or no family can be found for it
    def to_family(name):
        if a.rank != "family":
            return name
        if name.endswith("aceae"):
            return name
        for g in sorted(genera(name)):
            if g in family:
                return family[g]
        return family.get(name.split()[0], name)          # tribe labels such as Anthemideae

    # rolling their side up too, so both tables are counted at the same rank
    if a.rank == "family":
        for s in pub:
            rolled = collections.Counter()
            for t, v in pub[s].items():
                rolled[to_family(t)] += v
            pub[s] = rolled
        labels = {to_family(t) for t in labels}

    rep = {s: collections.Counter() for s in samples}
    translation = {}
    for r in rep_rows[1:]:
        label, how = translate(r[0]) if a.rank == "label" else (to_family(r[0]), "family roll-up")
        translation[r[0]] = (label, how)
        for s, v in zip(samples, r[1:]):
            rep[s][label] += float(v)

    shared = sorted(set(pub) & set(rep))
    with open(a.out, "w") as fh:
        fh.write(f"# sheet={a.sheet} samples_published={len(pub)} samples_replica={len(rep)} matched={len(shared)}\n")
        fh.write("sample\treads_published\treads_replica\tbray_curtis_similarity\tspearman_present_taxa\t"
                 "top_published\ttop_replica\ttop_agrees\n")
        bcs, rhos, agree = [], [], 0
        for s in shared:
            p, q = pub[s], rep[s]
            tp, tq = sum(p.values()) or 1, sum(q.values()) or 1
            # compare proportions, not counts: the two runs do not keep the same number of reads
            present = sorted({t for t in p if p[t] > 0} | {t for t in q if q[t] > 0})
            x = [p.get(t, 0) / tp for t in present]; y = [q.get(t, 0) / tq for t in present]
            bc = sum(min(u, v) for u, v in zip(x, y))          # = 1 - Bray-Curtis dissimilarity for proportions
            rho = spearman(x, y)
            top_p = max(p, key=p.get); top_q = max(q, key=q.get)
            agree += top_p == top_q; bcs.append(bc); rhos.append(rho)
            fh.write(f"{s}\t{int(tp)}\t{int(tq)}\t{bc:.3f}\t{rho:.3f}\t{top_p}\t{top_q}\t{top_p == top_q}\n")
        ok = [r for r in rhos if r == r]                       # r == r drops the nan rhos
        fh.write(f"# median_bray_curtis_similarity={statistics.median(bcs):.3f}  median_spearman={statistics.median(ok) if ok else float('nan'):.3f}  "
                 f"dominant_taxon_agrees={agree}/{len(shared)}\n")
    with open(a.out.replace(".tsv", ".name_translation.tsv"), "w") as fh:
        fh.write("replica_taxon\tpublished_label\thow\n")
        for k, (lab, how) in sorted(translation.items()):
            fh.write(f"{k}\t{lab}\t{how}\n")


if __name__ == "__main__":
    main()
